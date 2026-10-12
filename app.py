# -*- coding: utf-8 -*-
"""
استمارة تدقيق المعطيات الخاصة بالولي و التلميذ
الثانوية التأهيلية علال بن عبد الله — مديرية النواصر

Application Streamlit (interface en arabe, droite -> gauche) :
  - Le nom et le prénom (tuteur + élève) se saisissent en DEUX champs :
    un champ en arabe et un champ en français (comme sur les fiches Massar)
  - Date de naissance de l'élève obligatoire
  - Tous les champs sont obligatoires
  - Enregistrement automatique dans data/verifications.xlsx
  - Espace administrateur (barre latérale) protégé par mot de passe
    pour consulter et télécharger le fichier Excel.

Lancement :  streamlit run app.py
"""

import os
import re
from datetime import datetime, date

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# Configuration générale
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_DIR = os.path.join(BASE_DIR, "data")
EXCEL_PATH = os.path.join(EXCEL_DIR, "verifications.xlsx")

# Mot de passe de l'espace administrateur.
# Peut être changé via la variable d'environnement ADMIN_PASSWORD.
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin2026")

# ----------------------------------------------------------------------
# معلومات المؤسسة — Établissement
# ----------------------------------------------------------------------
SCHOOL_NAME = "الثانوية التأهيلية علال بن عبد الله"
DIRECTORATE = "مديرية النواصر"

# السنة الدراسية محسوبة تلقائيا (سبتمبر -> يوليوز من السنة الموالية)
_today = date.today()
_sy_start = _today.year if _today.month >= 7 else _today.year - 1
SCHOOL_YEAR = f"{_sy_start}/{_sy_start + 1}"

# ----------------------------------------------------------------------
# Noms de colonnes du fichier Excel (en arabe) — les noms / prénoms
# existent en deux versions : (بالعربية) et (بالفرنسية)
# ----------------------------------------------------------------------
COL_DATE = "تاريخ التسجيل"
COL_NOM_P_AR = "لقب ولي الأمر (بالعربية)"
COL_NOM_P_FR = "لقب ولي الأمر (بالفرنسية)"
COL_PRENOM_P_AR = "الاسم الشخصي لولي الأمر (بالعربية)"
COL_PRENOM_P_FR = "الاسم الشخصي لولي الأمر (بالفرنسية)"
COL_CIN = "رقم البطاقة الوطنية للتعريف"
COL_TEL = "رقم الهاتف"
COL_ADRESSE = "العنوان"
COL_NOM_E_AR = "لقب التلميذ(ة) (بالعربية)"
COL_NOM_E_FR = "لقب التلميذ(ة) (بالفرنسية)"
COL_PRENOM_E_AR = "الاسم الشخصي للتلميذ(ة) (بالعربية)"
COL_PRENOM_E_FR = "الاسم الشخصي للتلميذ(ة) (بالفرنسية)"
COL_DOB = "تاريخ ميلاد التلميذ(ة)"
COL_CODE = "رمز مسار"

COLUMNS = [
    COL_DATE,
    COL_NOM_P_AR, COL_NOM_P_FR, COL_PRENOM_P_AR, COL_PRENOM_P_FR,
    COL_CIN, COL_TEL, COL_ADRESSE,
    COL_NOM_E_AR, COL_NOM_E_FR, COL_PRENOM_E_AR, COL_PRENOM_E_FR,
    COL_DOB, COL_CODE,
]

FIELD_KEYS = [
    "parent_nom_ar", "parent_nom_fr", "parent_prenom_ar", "parent_prenom_fr",
    "cin", "tel", "adresse",
    "eleve_nom_ar", "eleve_nom_fr", "eleve_prenom_ar", "eleve_prenom_fr",
    "code_massar",
]

# ----------------------------------------------------------------------
# Page + style (arabe, droite -> gauche)
# ----------------------------------------------------------------------
st.set_page_config(
    page_title="استمارة تدقيق المعطيات — ثانوية علال بن عبد الله",
    layout="centered",
)

st.markdown(
    """
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700&display=swap');

        .stApp {
            font-family: 'Cairo', 'Segoe UI', 'Tahoma', sans-serif;
        }
        /* RTL posé sur les conteneurs de contenu et non sur .stApp :
           direction sur la racine casse le tiroir latéral de Streamlit
           en vue mobile (bande écrasée qui chevauche le formulaire). */
        .block-container,
        [data-testid="stMainBlockContainer"] {
            direction: rtl;
            text-align: right;
        }
        /* Vue bureau (>= 641 px) : RTL global, comme initialement
           (barre latérale à droite, mise en page inchangée). */
        @media (min-width: 641px) {
            .stApp {
                direction: rtl;
                text-align: right;
            }
        }
        .stTextInput input {
            direction: rtl;
            text-align: right;
            border-radius: 8px;
        }
        .stTextInput label, .stDateInput label {
            font-weight: 600;
            font-size: 0.93rem;
            line-height: 1.5;
        }
        /* Bouton principal d'envoi */
        div.stButton > button {
            width: 100%;
            background-color: #1F4E79;
            color: #ffffff;
            font-size: 1.05rem;
            font-weight: 700;
            border: none;
            border-radius: 8px;
            padding: 0.6rem 0;
        }
        div.stButton > button:hover {
            background-color: #163A5A;
            color: #ffffff;
        }
        div.stButton > button p {
            text-align: center !important;
        }
        /* Bouton de téléchargement (admin) */
        div.stDownloadButton > button {
            width: 100%;
            background-color: #2E7D32;
            color: #ffffff;
            font-weight: 700;
            border: none;
            border-radius: 8px;
            padding: 0.5rem 0;
        }
        div.stDownloadButton > button:hover {
            background-color: #1B5E20;
            color: #ffffff;
        }
        /* Titres de sections */
        .section-title {
            background-color: #EAF1F8;
            border-right: 5px solid #1F4E79;
            padding: 8px 14px;
            border-radius: 6px;
            font-weight: 700;
            font-size: 1.05rem;
            color: #1F4E79;
        }
        /* Les champs à contenu latin/chiffres passent en gauche-à-droite
           pour éviter l'inversion visuelle des groupes de caractères
           (sélecteurs par sous-chaîne de l'étiquette) */
        .stTextInput:has(input[aria-label*="الهاتف"]) input,
        .stTextInput:has(input[aria-label*="الوطنية"]) input,
        .stTextInput:has(input[aria-label*="مسار"]) input,
        .stTextInput:has(input[aria-label*="بالفرنسية"]) input {
            direction: ltr;
            text-align: right;
        }
        /* La table d'aperçu reste en gauche-à-droite (grille technique) */
        [data-testid="stDataFrame"] { direction: ltr; }
        /* Barre latérale : mécanique native (LTR). La direction RTL posée
           sur .stApp casse le tiroir latéral en vue mobile (bande écrasée
           qui chevauche le formulaire). Le contenu interne de la barre
           reste en RTL pour l'arabe. */
        [data-testid="stSidebar"],
        [data-testid="stSidebarCollapsedControl"] {
            direction: ltr;
        }
        [data-testid="stSidebar"] .block-container,
        [data-testid="stSidebarContent"] {
            direction: rtl;
            text-align: right;
        }
        /* Masquer les éléments techniques de Streamlit :
           le bouton Deploy, le menu ⋮ et le widget d'état sont masqués
           via .streamlit/config.toml (toolbarMode = "minimal") et non via
           CSS, car cacher [data-testid="stToolbar"] masque aussi le
           contrôle d'ouverture de la barre latérale sur téléphone. */
    </style>
    """,
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------
# Expressions régulières de validation
# ----------------------------------------------------------------------
# Noms en caractères arabes (lettres arabes + tatweel), espaces internes
_AR_LETTERS = r"\u0621-\u064A\u0640"
AR_NAME_RE = re.compile(rf"^[{_AR_LETTERS}]+(?:[\s{_AR_LETTERS}]+)*$")

# Noms en caractères latins (lettres + accents français), espaces,
# tirets et apostrophes internes
_FR_LETTERS = r"A-Za-z\u00C0-\u00D6\u00D8-\u00F6\u00F8-\u024F"
FR_NAME_RE = re.compile(rf"^[{_FR_LETTERS}]+(?:[\s'\-][{_FR_LETTERS}]+)*$")

CIN_RE = re.compile(r"^[A-Z]{1,2}\d{4,8}$")              # ex : A123456
PHONE_RE = re.compile(r"^(?:(?:\+|00)212|0)[1-7]\d{8}$")  # ex : 0612345678
CODE_RE = re.compile(r"^[A-Z]\d{8,12}$")                 # ex : R1300254789


def _clean_name(value: str) -> str:
    """Supprime les espaces superflus (début, fin, doubles espaces)."""
    return re.sub(r"\s+", " ", (value or "").strip())


# ----------------------------------------------------------------------
# Fonctions utilitaires (lecture / écriture Excel)
# ----------------------------------------------------------------------
def load_dataframe() -> pd.DataFrame:
    """Charge le fichier Excel existant, ou retourne un tableau vide."""
    if os.path.exists(EXCEL_PATH):
        try:
            # dtype=str : préserve les zéros initiaux (téléphone, codes...)
            df = pd.read_excel(EXCEL_PATH, engine="openpyxl", dtype=str)
            # Harmonise l'ordre des colonnes (compat anciens fichiers)
            return df.reindex(columns=COLUMNS)
        except Exception:
            pass
    return pd.DataFrame(columns=COLUMNS)


def _style_excel(path: str) -> None:
    """Mise en forme du classeur : en-tête stylé, largeurs, affichage RTL."""
    wb = load_workbook(path)
    ws = wb.active
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    widths = [17, 15, 15, 19, 19, 20, 14, 30, 15, 15, 19, 19, 17, 14]

    for idx, cell in enumerate(ws[1], start=1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        if idx <= len(widths):
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if cell.column in (7, 13, 14):
                # Téléphone + date de naissance + code Massar :
                # forcer le format texte
                cell.number_format = "@"

    ws.freeze_panes = "A2"
    wb.save(path)


def save_dataframe(df: pd.DataFrame) -> None:
    """Enregistre le tableau dans le fichier Excel + mise en forme."""
    os.makedirs(EXCEL_DIR, exist_ok=True)
    df.to_excel(EXCEL_PATH, index=False, engine="openpyxl")
    _style_excel(EXCEL_PATH)


def append_row(row: dict) -> None:
    """Ajoute une vérification au fichier Excel."""
    df = load_dataframe()
    df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
    save_dataframe(df)


def code_massar_exists(code: str) -> bool:
    """Vérifie si le code Massar est déjà enregistré."""
    df = load_dataframe()
    if df.empty or COL_CODE not in df.columns:
        return False
    existing = df[COL_CODE].astype(str).str.strip().str.upper()
    return code.strip().upper() in existing.values


# ----------------------------------------------------------------------
# Validation du formulaire
# ----------------------------------------------------------------------
def validate_inputs(parent_nom_ar, parent_prenom_ar,
                    parent_nom_fr, parent_prenom_fr,
                    cin, tel, adresse,
                    eleve_nom_ar, eleve_prenom_ar,
                    eleve_nom_fr, eleve_prenom_fr,
                    eleve_dob, code_massar):
    """Vérifie tous les champs (obligatoires + script arabe/français +
    formats) et normalise le CIN, le téléphone et le code Massar.
    Retourne (erreurs, valeurs)."""
    errors = []

    nom_p_ar = _clean_name(parent_nom_ar)
    prenom_p_ar = _clean_name(parent_prenom_ar)
    nom_p_fr = _clean_name(parent_nom_fr)
    prenom_p_fr = _clean_name(parent_prenom_fr)
    nom_e_ar = _clean_name(eleve_nom_ar)
    prenom_e_ar = _clean_name(eleve_prenom_ar)
    nom_e_fr = _clean_name(eleve_nom_fr)
    prenom_e_fr = _clean_name(eleve_prenom_fr)

    # --- Tuteur : noms en arabe ---------------------------------------
    if not nom_p_ar:
        errors.append("المرجو إدخال لقب ولي الأمر بالحروف العربية.")
    elif not AR_NAME_RE.match(nom_p_ar):
        errors.append("لقب ولي الأمر: المرجو الكتابة بالحروف العربية فقط "
                      "(بدون أرقام أو حروف لاتينية).")

    if not prenom_p_ar:
        errors.append("المرجو إدخال الاسم الشخصي لولي الأمر بالحروف العربية.")
    elif not AR_NAME_RE.match(prenom_p_ar):
        errors.append("الاسم الشخصي لولي الأمر: المرجو الكتابة بالحروف "
                      "العربية فقط (بدون أرقام أو حروف لاتينية).")

    # --- Tuteur : noms en français ------------------------------------
    if not nom_p_fr:
        errors.append("المرجو إدخال لقب ولي الأمر بالحروف الفرنسية.")
    elif not FR_NAME_RE.match(nom_p_fr):
        errors.append("لقب ولي الأمر (بالفرنسية): المرجو الكتابة بالحروف "
                      "الفرنسية فقط (بدون حروف عربية أو أرقام).")

    if not prenom_p_fr:
        errors.append("المرجو إدخال الاسم الشخصي لولي الأمر بالحروف الفرنسية.")
    elif not FR_NAME_RE.match(prenom_p_fr):
        errors.append("الاسم الشخصي لولي الأمر (بالفرنسية): المرجو الكتابة "
                      "بالحروف الفرنسية فقط (بدون حروف عربية أو أرقام).")

    # --- CIN ------------------------------------------------------------
    cin_clean = cin.strip().upper()
    if not cin_clean:
        errors.append("المرجو إدخال رقم البطاقة الوطنية للتعريف.")
    elif not CIN_RE.match(cin_clean):
        errors.append("رقم البطاقة الوطنية للتعريف غير صحيح (مثال: A123456).")

    # --- Téléphone ------------------------------------------------------
    tel_clean = re.sub(r"[\s\-.()]", "", tel.strip())
    if not tel_clean:
        errors.append("المرجو إدخال رقم الهاتف.")
    elif not PHONE_RE.match(tel_clean):
        errors.append("رقم الهاتف غير صحيح (مثال: 0612345678).")

    # --- Adresse --------------------------------------------------------
    if len(adresse.strip()) < 5:
        errors.append("المرجو إدخال العنوان الكامل.")

    # --- Élève : noms en arabe -----------------------------------------
    if not nom_e_ar:
        errors.append("المرجو إدخال لقب التلميذ(ة) بالحروف العربية.")
    elif not AR_NAME_RE.match(nom_e_ar):
        errors.append("لقب التلميذ(ة): المرجو الكتابة بالحروف العربية فقط "
                      "(بدون أرقام أو حروف لاتينية).")

    if not prenom_e_ar:
        errors.append("المرجو إدخال الاسم الشخصي للتلميذ(ة) بالحروف العربية.")
    elif not AR_NAME_RE.match(prenom_e_ar):
        errors.append("الاسم الشخصي للتلميذ(ة): المرجو الكتابة بالحروف "
                      "العربية فقط (بدون أرقام أو حروف لاتينية).")

    # --- Élève : noms en français ---------------------------------------
    if not nom_e_fr:
        errors.append("المرجو إدخال لقب التلميذ(ة) بالحروف الفرنسية.")
    elif not FR_NAME_RE.match(nom_e_fr):
        errors.append("لقب التلميذ(ة) (بالفرنسية): المرجو الكتابة بالحروف "
                      "الفرنسية فقط (بدون حروف عربية أو أرقام).")

    if not prenom_e_fr:
        errors.append("المرجو إدخال الاسم الشخصي للتلميذ(ة) بالحروف الفرنسية.")
    elif not FR_NAME_RE.match(prenom_e_fr):
        errors.append("الاسم الشخصي للتلميذ(ة) (بالفرنسية): المرجو الكتابة "
                      "بالحروف الفرنسية فقط (بدون حروف عربية أو أرقام).")

    # --- Date de naissance de l'élève -----------------------------------
    if eleve_dob is None:
        errors.append("المرجو تحديد تاريخ ميلاد التلميذ(ة).")
    dob_str = eleve_dob.strftime("%d/%m/%Y") if eleve_dob else ""

    # --- Code Massar ------------------------------------------------------
    code_clean = code_massar.strip().upper()
    if not code_clean:
        errors.append("المرجو إدخال رمز مسار.")
    elif not CODE_RE.match(code_clean):
        errors.append("رمز مسار غير صحيح (مثال: R1300254789).")

    values = {
        "nom_p_ar": nom_p_ar, "nom_p_fr": nom_p_fr,
        "prenom_p_ar": prenom_p_ar, "prenom_p_fr": prenom_p_fr,
        "cin": cin_clean, "tel": tel_clean, "adresse": adresse.strip(),
        "nom_e_ar": nom_e_ar, "nom_e_fr": nom_e_fr,
        "prenom_e_ar": prenom_e_ar, "prenom_e_fr": prenom_e_fr,
        "dob": dob_str, "code": code_clean,
    }
    return errors, values


# ----------------------------------------------------------------------
# Réinitialisation des champs — à exécuter AVANT l'instanciation des
# widgets pour que la valeur vide soit bien répercutée côté navigateur.
# ----------------------------------------------------------------------
if st.session_state.get("_clear_fields"):
    for _k in FIELD_KEYS:
        st.session_state[_k] = ""
    st.session_state["eleve_dob"] = None
    st.session_state.pop("_clear_fields")

if st.session_state.get("_logout"):
    st.session_state["admin_pwd"] = ""
    st.session_state["is_admin"] = False
    st.session_state.pop("_logout")

# ----------------------------------------------------------------------
# ترويسة المؤسسة + عنوان الاستمارة (بالعربية فقط)
# ----------------------------------------------------------------------
st.markdown(
    f"""
    <div style="text-align:center; border:2px solid #1F4E79; border-radius:10px;
                padding:12px 10px 8px; margin-bottom:10px; background:#F7FAFD;">
        <p style="margin:0; color:#555; font-size:0.95em;">{DIRECTORATE}</p>
        <p style="margin:4px 0 0; color:#1F4E79; font-weight:700; font-size:1.2em;">{SCHOOL_NAME}</p>
        <p style="margin:4px 0 0; color:#777; font-size:0.85em;">السنة الدراسية: {SCHOOL_YEAR}</p>
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    '<h1 style="text-align:center; color:#1F4E79; margin-bottom:0; '
    'font-size:1.7rem; line-height:1.5;">'
    'استمارة تدقيق المعطيات الخاصة بالولي و التلميذ</h1>',
    unsafe_allow_html=True,
)
st.markdown(
    '<p style="text-align:center; color:#555;">'
    'المرجو من ولي الأمر تدقيق المعطيات الخاصة به وللتلميذ(ة) '
    'وتعبئة جميع الخانات بدقة، فجميع الحقول إلزامية.<br>'
    'كما يُرجى إدخال اللقب والاسم الشخصي بالعربية وبالفرنسية.</p>',
    unsafe_allow_html=True,
)
st.markdown("---")

# Message de succès affiché après un envoi réussi
if "success_msg" in st.session_state:
    st.success(st.session_state["success_msg"])
    if "success_details" in st.session_state:
        st.info(st.session_state["success_details"])
    st.markdown(
        '<p style="color:#2E7D32;">'
        'يمكنكم تدقيق معطيات تلميذ(ة) آخر عبر تعبئة الاستمارة من جديد.</p>',
        unsafe_allow_html=True,
    )

st.markdown('<div class="section-title">أولا: معلومات ولي الأمر</div>',
            unsafe_allow_html=True)

# Nom : un champ en arabe + un champ en français
c1, c2 = st.columns(2)
with c1:
    parent_nom_ar = st.text_input(
        "اللقب (بالعربية) *", placeholder="مثال: العلمي",
        key="parent_nom_ar")
with c2:
    parent_nom_fr = st.text_input(
        "اللقب (بالفرنسية) *", placeholder="مثال: Alami",
        key="parent_nom_fr")

# Prénom : un champ en arabe + un champ en français
c3, c4 = st.columns(2)
with c3:
    parent_prenom_ar = st.text_input(
        "الاسم الشخصي (بالعربية) *", placeholder="مثال: محمد",
        key="parent_prenom_ar")
with c4:
    parent_prenom_fr = st.text_input(
        "الاسم الشخصي (بالفرنسية) *", placeholder="مثال: Mohammed",
        key="parent_prenom_fr")

c5, c6 = st.columns(2)
with c5:
    cin = st.text_input("رقم البطاقة الوطنية للتعريف *",
                        placeholder="A123456", max_chars=12, key="cin")
with c6:
    tel = st.text_input("رقم الهاتف *",
                        placeholder="0612345678", max_chars=16, key="tel")

adresse = st.text_input("العنوان الكامل *",
                        placeholder="مثال: حي النهضة، الدار البيضاء",
                        key="adresse")

st.markdown('<div class="section-title" style="margin-top:12px;">'
            "ثانيا: معلومات التلميذ(ة)</div>",
            unsafe_allow_html=True)

# Nom de l'élève : un champ en arabe + un champ en français
c7, c8 = st.columns(2)
with c7:
    eleve_nom_ar = st.text_input(
        "لقب التلميذ(ة) (بالعربية) *", placeholder="مثال: العلمي",
        key="eleve_nom_ar")
with c8:
    eleve_nom_fr = st.text_input(
        "لقب التلميذ(ة) (بالفرنسية) *", placeholder="مثال: Alami",
        key="eleve_nom_fr")

# Prénom de l'élève : un champ en arabe + un champ en français
c9, c10 = st.columns(2)
with c9:
    eleve_prenom_ar = st.text_input(
        "الاسم الشخصي للتلميذ(ة) (بالعربية) *", placeholder="مثال: يوسف",
        key="eleve_prenom_ar")
with c10:
    eleve_prenom_fr = st.text_input(
        "الاسم الشخصي للتلميذ(ة) (بالفرنسية) *", placeholder="مثال: Youssef",
        key="eleve_prenom_fr")

c11, c12 = st.columns(2)
with c11:
    eleve_dob = st.date_input(
        "تاريخ ميلاد التلميذ(ة) *",
        value=None,
        min_value=date(1980, 1, 1),
        max_value=date(2020, 12, 31),
        format="DD/MM/YYYY",
        key="eleve_dob",
    )
with c12:
    code_massar = st.text_input("رمز مسار *",
                                placeholder="R1300254789",
                                max_chars=13, key="code_massar")

st.markdown('<p style="color:#888; font-size:0.85em;">'
            '(*) جميع الحقول إلزامية — يُدوَّن اللقب والاسم الشخصي '
            'بالعربية وبالفرنسية.</p>',
            unsafe_allow_html=True)

submitted = st.button("تسليم الاستمارة", key="submit_btn")

# ----------------------------------------------------------------------
# Traitement de la soumission
# ----------------------------------------------------------------------
if submitted:
    errors, values = validate_inputs(
        parent_nom_ar, parent_prenom_ar, parent_nom_fr, parent_prenom_fr,
        cin, tel, adresse,
        eleve_nom_ar, eleve_prenom_ar, eleve_nom_fr, eleve_prenom_fr,
        eleve_dob, code_massar,
    )

    if errors:
        for err in errors:
            st.error(err)
    elif code_massar_exists(values["code"]):
        st.error("رمز مسار مسجل مسبقا. المرجو التحقق من الرمز "
                 "أو الاتصال بإدارة المؤسسة.")
    else:
        row = {
            COL_DATE: datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            COL_NOM_P_AR: values["nom_p_ar"],
            COL_NOM_P_FR: values["nom_p_fr"],
            COL_PRENOM_P_AR: values["prenom_p_ar"],
            COL_PRENOM_P_FR: values["prenom_p_fr"],
            COL_CIN: values["cin"],
            COL_TEL: values["tel"],
            COL_ADRESSE: values["adresse"],
            COL_NOM_E_AR: values["nom_e_ar"],
            COL_NOM_E_FR: values["nom_e_fr"],
            COL_PRENOM_E_AR: values["prenom_e_ar"],
            COL_PRENOM_E_FR: values["prenom_e_fr"],
            COL_DOB: values["dob"],
            COL_CODE: values["code"],
        }
        try:
            append_row(row)
            st.session_state["success_msg"] = (
                "تم تسليم الاستمارة بنجاح. شكرا لتعاونكم!"
            )
            st.session_state["success_details"] = (
                f"التلميذ(ة): {values['prenom_e_ar']} {values['nom_e_ar']} "
                f"({values['prenom_e_fr']} {values['nom_e_fr']})"
                f" — رمز مسار: {values['code']}"
                f" — تاريخ الميلاد: {values['dob']}"
            )
            # Vider les champs pour une nouvelle inscription
            # (traité en début de script suivant, avant l'instanciation
            # des widgets, pour que le navigateur efface aussi l'affichage)
            st.session_state["_clear_fields"] = True
            st.rerun()
        except Exception:
            st.error("حدث خطأ أثناء حفظ البيانات. المرجو إغلاق ملف Excel "
                     "إذا كان مفتوحا والمحاولة من جديد.")

st.markdown("---")
st.markdown(
    '<p style="text-align:center; color:#777; font-size:0.85em;">'
    'جميع المعلومات المدخلة سرية وتستعمل حصريا للأغراض الإدارية '
    'للمؤسسة.</p>',
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------
# Espace administrateur — فضاء المشرف (barre latérale)
# ----------------------------------------------------------------------
with st.sidebar:
    st.markdown('<div class="section-title">فضاء المشرف</div>',
                unsafe_allow_html=True)
    st.caption("فضاء مخصص لإدارة المؤسسة لمعاينة وتحميل ملف التدقيقات.")

    pwd = st.text_input("كلمة السر", type="password", key="admin_pwd")

    if pwd:
        if pwd == ADMIN_PASSWORD:
            st.session_state["is_admin"] = True
        else:
            st.session_state["is_admin"] = False
            st.error("كلمة السر غير صحيحة.")
    elif st.session_state.get("is_admin"):
        st.session_state["is_admin"] = False

    if st.session_state.get("is_admin"):
        df_admin = load_dataframe()
        n = len(df_admin)

        st.markdown(
            f'<div style="background:#E8F5E9; padding:10px; border-radius:8px;'
            f' text-align:center; font-weight:700; color:#2E7D32;">'
            f'عدد التدقيقات: {n}</div>',
            unsafe_allow_html=True,
        )

        if os.path.exists(EXCEL_PATH) and n > 0:
            st.markdown("##### معاينة البيانات")
            st.dataframe(df_admin, height=250, hide_index=True)

            with open(EXCEL_PATH, "rb") as f:
                excel_bytes = f.read()
            st.download_button(
                label="تحميل ملف Excel",
                data=excel_bytes,
                file_name=f"تدقيق_المعطيات_ثانوية_علال_بن_عبد_الله_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="download_btn",
            )
        else:
            st.info("لا توجد أي تدقيقات بعد.")

        if st.button("تحديث البيانات", key="refresh_btn"):
            st.rerun()

        if st.button("خروج", key="logout_btn"):
            st.session_state["_logout"] = True
            st.rerun()
